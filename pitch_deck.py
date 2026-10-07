from pptx import Presentation

presentation = Presentation()

slide = presentation.slides.add_slide(
    presentation.slide_layouts[0]
)

slide.shapes.title.text = "From Skills to Success"

presentation.save("pitch_deck.pptx")

print("Pitch deck created successfully!")